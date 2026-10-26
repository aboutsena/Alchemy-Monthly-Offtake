"""Independent synthetic transaction and preservation tests for portable stage."""
from __future__ import annotations

import copy
import importlib.util
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import unittest
import zipfile

PROGRAM = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROGRAM))
import portable_stage as stage

ENGINE = Path(os.environ.get("ALCHEMY_TEST_ENGINE", PROGRAM / "engine/Automacao off-take"))


def fixture(directory, basis="ex_vat"):
    """Author a tiny independent workbook, not copied commercial workbooks."""
    import openpyxl
    directory = Path(directory)
    source_dir = directory / "inputs"
    source_dir.mkdir()
    source = source_dir / "SYNTHETIC TEST ONLY.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for number in range(1, 13):
        sheet = workbook.create_sheet(f"{number:02}")
        sheet["C2"] = "Synthetic Outlet"
        for address, value in {"B13": "Code", "G13": "Product", "H13": "Size", "J13": "Wholesaler",
                               "K13": "Price Inc. VAT", "L13": "Price Ex. VAT", "N13": "Vol. (Btls.)",
                               "P13": "Rebate (%)", "B14": "SYNTHETIC-SKU", "G14": "Synthetic Gin",
                               "H14": "700ml", "P14": .1}.items():
            sheet[address] = value
        for column in "NRST":
            sheet[column + "4"] = "=SUM(" + column + "14:" + column + "14)"
        sheet["R14"] = "=IF(K14>0,K14*N14,L14*1.07*N14)"
        sheet["S14"] = "=ROUND(R14/1.07,2)"
        sheet["T14"] = "=P14*S14"
        sheet["U14"] = '=IF(AND(K14>0,L14>0),"Input Price only 1 column","OK")'
    workbook.create_sheet("SUMMARY")["B1"] = "='10'!S4"
    workbook.save(source)
    workbook.close()
    # Explicit Excel-compatible caches are fixture data. These zeros describe an
    # empty template, never sales or expected outputs of the code under test.
    with zipfile.ZipFile(source) as archive:
        original = [(copy.copy(info), archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(source, "w") as archive:
        for info, data in original:
            if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                text = data.decode("utf-8")
                def cached(match):
                    text = match.group()
                    if "<f" not in text:
                        return text
                    address = re.search(r'\br="([A-Z]+\d+)"', text)[1]
                    text = re.sub(r"<v(?:\s*/>|>.*?</v>)", "", text)
                    text = re.sub(r'\s+t="[^"]*"', "", text, count=1)
                    if address.startswith("U"):
                        text = text.replace("<c ", '<c t="str" ', 1)
                        return text.replace("</c>", "<v>OK</v></c>")
                    return text.replace("</c>", "<v>0</v></c>")
                data = re.sub(r"<c\b[^>]*>.*?</c>", cached, text).encode("utf-8")
            archive.writestr(info, data)
    report = directory / "SYNTHETIC SOURCE ONLY.json"
    report.write_text('{"notice":"SYNTHETIC TEST ONLY"}', encoding="utf-8")
    engine = stage.engine_modules(ENGINE)
    model, _ = engine.components("10")
    with zipfile.ZipFile(source) as archive:
        month = model.c.make_sheet_spec(source, archive, "10", "outlet")
        dependent = model.c.make_sheet_spec(source, archive, "SUMMARY", "dependency", {
            "cache_addresses": ["B1"], "expected_values": {"B1": 200}, "reason": "Synthetic current-month dependency"})
    price = 100 if basis == "ex_vat" else 107
    edits = {"J14": "Synthetic Supplier", "K14": price if basis == "inc_vat" else None,
             "L14": price if basis == "ex_vat" else None, "N14": 2}
    month.update(month="10", clear_cells=[], edits=edits, mapping_additions=[], parameter_overrides=[], formula_overrides=[])
    values = {"N4": 2, "R4": 214, "S4": 200, "T4": 20, "R14": 214,
              "S14": 200, "T14": 20, "U14": "OK", **edits}
    month["expected_values"].update(values)
    month["expected_totals"] = {"N": 2, "R": 214, "S": 200, "T": 20}
    dependent["edits"] = {}
    record = {"record_id": "synthetic-only-row-1", "period": ["2026-10-01", "2026-10-31"],
              "source_period_verified": True, "dates": ["2026-10-15"], "source": str(report),
              "source_sha256": stage.sha(report), "quantity_bottles": 2, "amount": price * 2,
              "price_basis": basis, "unit_price": price}
    row = {"row": 14, "qty": 2, "price": price, "price_decimal": str(price), "rebate_rate": .1,
           "price_basis": basis, "wholesaler": "Synthetic Supplier", "amount": price * 2,
           "expected_R": 214, "expected_S": 200, "expected_T": 20, "source_records": [record]}
    return {"id": "job-001", "kind": "outlet", "outlet": "Synthetic Outlet", "path": str(source),
            "output_name": source.name, "sha256": stage.sha(source), "month": "10", "period": "2026-10",
            "dry_run": True, "sheets": [month, dependent], "rows": [row],
            "totals": {"bottles": 2, "inc_vat": 214, "ex_vat": 200, "rebate": 20}}


@unittest.skipUnless(ENGINE.is_dir(), "Reviewed engine distribution is required")
class PortableStageTests(unittest.TestCase):
    def setUp(self):
        test_root = PROGRAM.parent / "test-work"
        test_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="alchemy-stage-synthetic-", dir=test_root)
        self.assertTrue(Path(self.temp.name).resolve().is_relative_to(test_root.resolve()))
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.job = fixture(self.root)
        self.original_sha = stage.sha(self.job["path"])

    def run_job(self, job=None, staging=None):
        return stage.calculate_job(job or self.job, ENGINE, staging or self.root / "scratch")

    def blocked(self, mutate):
        mutate(self.job)
        with self.assertRaises((ValueError, AssertionError, KeyError)):
            self.run_job()
        self.assertEqual(stage.sha(self.job["path"]), self.original_sha)
        self.assertFalse((self.root / "scratch/job-001-stage-manifest.json").exists())

    def test_ex_vat_fresh_formulas_and_package_preservation(self):
        proof = self.run_job()
        self.assertEqual(proof["result"], "passed")
        self.assertFalse(proof["independent_second_formula_engine"])
        self.assertEqual(proof["fresh_formula_cache_checks"], 9)
        self.assertEqual(proof["independent_decimal_row_checks"], 1)
        self.assertEqual(stage.sha(self.job["path"]), self.original_sha)
        with zipfile.ZipFile(self.job["path"]) as original, zipfile.ZipFile(proof["path"]) as output:
            self.assertEqual(original.namelist(), output.namelist())
            allowed = {s["sheet_part"] for s in self.job["sheets"]}
            for part in original.namelist():
                if part not in allowed:
                    self.assertEqual(original.read(part), output.read(part), part)
            before = original.read(self.job["sheets"][0]["sheet_part"]).decode("utf-8")
            after = output.read(self.job["sheets"][0]["sheet_part"]).decode("utf-8")
            self.assertEqual(re.findall(r"<f\b[^>]*>.*?</f>", before), re.findall(r"<f\b[^>]*>.*?</f>", after))
            self.assertEqual(re.search(r'<c\b[^>]*\br="P14"[^>]*>.*?</c>', before).group(),
                             re.search(r'<c\b[^>]*\br="P14"[^>]*>.*?</c>', after).group())

    def test_inc_vat_explicit_basis_passes(self):
        second = self.root / "inc"
        second.mkdir()
        job = fixture(second, "inc_vat")
        self.assertEqual(self.run_job(job)["result"], "passed")

    def test_rejects_changed_source_hash(self):
        self.blocked(lambda j: j.update(sha256="0" * 64))

    def test_rejects_month_period_disagreement(self):
        self.blocked(lambda j: j.update(month="09"))

    def test_rejects_wrong_transaction_month(self):
        self.blocked(lambda j: j["rows"][0]["source_records"][0].update(dates=["2026-09-15"]))

    def test_rejects_unknown_source_period(self):
        self.blocked(lambda j: j["rows"][0]["source_records"][0].update(source_period_verified=False))

    def test_rejects_duplicate_occurrence(self):
        self.blocked(lambda j: j["rows"][0]["source_records"].append(copy.deepcopy(j["rows"][0]["source_records"][0])))

    def test_rejects_missing_source_provenance(self):
        self.blocked(lambda j: j["rows"][0].update(source_records=[]))

    def test_rejects_quantity_tampering(self):
        self.blocked(lambda j: j["sheets"][0]["edits"].update(N14=3))

    def test_rejects_unknown_vat_basis(self):
        self.blocked(lambda j: j["rows"][0].update(price_basis="unknown"))

    def test_rejects_price_tampering(self):
        self.blocked(lambda j: j["sheets"][0]["edits"].update(L14=150))

    def test_rejects_two_price_columns(self):
        self.blocked(lambda j: j["sheets"][0]["edits"].update(K14=107))

    def test_rejects_changed_rebate_rate(self):
        self.blocked(lambda j: j["rows"][0].update(rebate_rate=.2))

    def test_rejects_parameter_override(self):
        self.blocked(lambda j: j["sheets"][0].update(parameter_overrides=[{"address": "P14", "value": .2}]))

    def test_rejects_formula_override(self):
        self.blocked(lambda j: j["sheets"][0].update(formula_overrides=[{"address": "T14", "new_formula": "0"}]))

    def test_rejects_tampered_prepared_cache(self):
        self.blocked(lambda j: j["sheets"][0]["expected_values"].update(S14=201))

    def test_rejects_tampered_sheet_binding(self):
        self.blocked(lambda j: j["sheets"][0].update(original_sheet_sha256="0" * 64))

    def test_rejects_tampered_declared_total(self):
        self.blocked(lambda j: j["totals"].update(ex_vat=201))

    def test_rejects_unplanned_edits(self):
        self.blocked(lambda j: j["sheets"][0]["edits"].update(J15="unexpected"))

    def test_rejects_other_month_edit(self):
        self.blocked(lambda j: j["sheets"][1].update(name="09", edits={"N14": 2}))

    def test_rejects_output_traversal(self):
        self.blocked(lambda j: j.update(output_name="../outside.xlsx"))

    def test_rejects_output_over_inputs(self):
        with self.assertRaises(ValueError):
            self.run_job(staging=Path(self.job["path"]).parent)
        self.assertEqual(stage.sha(self.job["path"]), self.original_sha)

    def test_rejects_existing_output(self):
        target = self.root / "scratch"
        target.mkdir()
        existing = target / self.job["output_name"]
        existing.write_bytes(b"do not replace")
        with self.assertRaises(ValueError):
            self.run_job()
        self.assertEqual(existing.read_bytes(), b"do not replace")

    def test_rejects_existing_proof(self):
        target = self.root / "scratch"
        target.mkdir()
        proof = target / "job-001-stage-manifest.json"
        proof.write_text("do not replace", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.run_job()
        self.assertEqual(proof.read_text(encoding="utf-8"), "do not replace")


class NumberValidationTests(unittest.TestCase):
    def test_nonfinite_numbers_fail(self):
        for value in (float("inf"), float("nan"), True, None):
            with self.subTest(value=value), self.assertRaises(ValueError):
                stage.decimal(value)

    def test_blank_is_not_zero(self):
        self.assertFalse(stage.value_matches(None, 0))
        self.assertFalse(stage.value_matches("", 0))
        self.assertTrue(stage.value_matches(None, ""))


if __name__ == "__main__":
    unittest.main()
